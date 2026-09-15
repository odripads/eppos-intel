#!/usr/bin/env python3
"""xlsx → JSON build for the EPPOS incident map. Run at build time; the browser only ever fetches /data/*.json.

Usage: python3 scripts/build_data.py [eppos-media-intake.xlsx] [--wikidata path/to/wikidata_id_regions.json]

Writes data/periode.json, data/insiden.json, data/kasus_resmi.json, data/koordinat.json,
data/koordinat_gagal.json (places that could not be resolved), data/validasi.json (rejected rows),
data/manifest.json. Sheets other than periode/insiden/kasus_resmi are not read.

Rules enforced here (from PROMPT + PROJECT-SPEC v2):
  * a row without a source URL is a data bug: reported in validasi.json and NOT written to the output;
  * blank stays blank (null); nothing is inferred or filled;
  * derived titles stay marked `dari URL` unless data/judul_terverifikasi.json (from verify_titles.py) has the real <title>;
  * all rows including status_kurasi=keluar are written — the page filters; exclusion is transparent.
"""
from __future__ import annotations
import argparse, datetime as dt, hashlib, json, re, sys, unicodedata
from collections import Counter
from pathlib import Path
import openpyxl

ROOT = Path(__file__).resolve().parent.parent
DATA = ROOT / "data"
HEADER_ROW, HINT_ROW, FIRST_DATA_ROW = 2, 3, 4

INSIDEN_COLS = ["insiden_id", "tanggal", "provinsi", "kab_kota", "pelaku_jabatan", "sasaran_jenis", "mekanisme",
                "ringkasan_satu_kalimat", "hasil", "sumber_1_url", "sumber_1_outlet", "sumber_2_url", "status_verifikasi",
                "status_kurasi", "periode_pilpres", "judul_sumber_1", "judul_status", "diisi_oleh", "tanggal_isi"]
KASUS_COLS = ["kasus_id", "lembaga", "tahun", "provinsi", "kab_kota", "jenis_pelanggaran", "jumlah_kasus", "sumber_url",
              "periode_pilpres", "judul_sumber_1", "judul_status", "diisi_oleh", "tanggal_isi"]
URL_RE = re.compile(r"^https?://\S+$", re.I)

# ---- projection of the mockup (equirectangular into 1000x420) — verified below before anything is written
LON_MIN, LON_MAX, LAT_MIN, LAT_MAX = 95.195138, 141.0009, -10.91942, 5.877151
S = 1000 / (LON_MAX - LON_MIN)
OY = (420 - (LAT_MAX - LAT_MIN) * S) / 2
def project(lon, lat): return (lon - LON_MIN) * S, 420 - ((lat - LAT_MIN) * S + OY)

def check_projection():
    for name, lon, lat, ex, ey in (("Kota Makassar", 119.4144, -5.1477, 528.7, 267.3), ("Kupang", 123.5833, -10.1718, 619.8, 377.0)):
        x, y = project(lon, lat)
        if abs(x - ex) > 0.3 or abs(y - ey) > 0.3:
            sys.exit(f"PROYEKSI SALAH: {name} → ({x:.1f}, {y:.1f}), seharusnya ({ex}, {ey})")
    print(f"proyeksi OK  (s={S:.4f}, oy={OY:.2f}; Makassar → {project(119.4144, -5.1477)[0]:.1f},{project(119.4144, -5.1477)[1]:.1f})")

# ---- cell normalisation: never infer, only serialise
def cell(v):
    if v is None: return None
    if isinstance(v, dt.datetime): return v.date().isoformat()
    if isinstance(v, dt.date): return v.isoformat()
    if isinstance(v, float) and v.is_integer(): return int(v)
    s = str(v).strip()
    if s == "": return None
    return int(s) if re.fullmatch(r"\d{1,9}", s) else s   # digit-only text cells (tahun, jumlah_kasus) serialise as numbers; nothing else is coerced

def read_sheet(wb, name, first_row=FIRST_DATA_ROW):
    ws = wb[name]
    headers = [cell(c.value) for c in ws[HEADER_ROW]]
    rows = []
    for r in ws.iter_rows(min_row=first_row, values_only=False):
        vals = [cell(c.value) for c in r]
        if not any(v is not None for v in vals): continue
        rec = {h: v for h, v in zip(headers, vals) if h}
        rec["_row"] = r[0].row
        rows.append(rec)
    return headers, rows

def snake(s): return re.sub(r"[^a-z0-9]+", "_", s.lower()).strip("_")

# ---- periode: tab definitions come from the sheet, verbatim
def build_periode(wb):
    headers, rows = read_sheet(wb, "periode", first_row=HEADER_ROW + 1)  # periode has no hint row: data starts at row 3
    out = []
    for r in rows:
        if not r.get("periode") or not str(r["periode"]).startswith("Periode"): continue  # notes block below the table
        out.append({snake(k): v for k, v in r.items() if k != "_row"})
    return out

# ---- koordinat
def norm(s):
    s = unicodedata.normalize("NFKD", s).encode("ascii", "ignore").decode().lower()
    return re.sub(r"\s+", " ", s).strip()

PROV_ALIAS = {"di yogyakarta": "yogyakarta", "dki jakarta": "jakarta"}   # Wikidata labels the provinces plainly
# Hand-filled centroids for places Wikidata does not type as kab/kota (Jakarta's administrative cities). Source recorded per entry.
MANUAL = {("dki jakarta", "jakarta utara"): {"lat": -6.1385, "lon": 106.8637, "sumber": "manual: id.wikipedia.org/wiki/Jakarta_Utara (koordinat infobox)"}}
KAB_PREFIX = re.compile(r"^(kabupaten|kab\.?|kota administrasi|kota adm\.?|kota)\s+", re.I)

def load_wikidata(path):
    if not path or not Path(path).exists(): return None
    rows = json.load(open(path))
    idx = {"kab": {}, "kota": {}, "prov": {}}
    for r in rows:
        t = r["type"].lower(); lab = norm(r["label"])
        if "provin" in t: idx["prov"].setdefault(norm(KAB_PREFIX.sub("", lab)), []).append(r)
        elif "kota" in t or "city" in t: idx["kota"].setdefault(norm(KAB_PREFIX.sub("", lab)), []).append(r)
        else: idx["kab"].setdefault(norm(KAB_PREFIX.sub("", lab)), []).append(r)
    return idx

def resolve(prov, kab, wd, gagal):
    """Return a koordinat entry or None. Never guesses: multi-place strings fall back to the province centroid, flagged."""
    def src(r): return f"wikidata:{r['qid']} ({r['label']})"
    def prov_entry(reason):
        if not prov: return None
        key = PROV_ALIAS.get(norm(prov), norm(prov))
        hits = wd["prov"].get(key, []) if wd else []
        if len(hits) != 1:
            gagal.append({"provinsi": prov, "kab_kota": kab, "alasan": f"centroid provinsi tidak ditemukan/ambigu ({len(hits)} kandidat)"}); return None
        r = hits[0]
        return {"lat": r["lat"], "lon": r["lon"], "tingkat": "provinsi", "alasan": reason, "sumber": src(r)}
    if not prov and not kab:
        gagal.append({"provinsi": prov, "kab_kota": kab, "alasan": "baris nasional: tidak dipetakan"}); return None
    if not kab: return prov_entry("kab_kota kosong: agregat provinsi")
    if "," in kab or ":" in kab or re.search(r"\b\d+\b", kab):
        e = prov_entry(f"multi kab/kota, ditempatkan di centroid provinsi: “{kab}”")
        gagal.append({"provinsi": prov, "kab_kota": kab, "alasan": "beberapa kab/kota dalam satu sel; hanya centroid provinsi"}); return e
    if (norm(prov or ""), norm(kab)) in MANUAL:
        m = MANUAL[(norm(prov or ""), norm(kab))]; return {"lat": m["lat"], "lon": m["lon"], "tingkat": "kab_kota", "alasan": None, "sumber": m["sumber"]}
    if not wd:
        gagal.append({"provinsi": prov, "kab_kota": kab, "alasan": "berkas centroid Wikidata tidak tersedia"}); return None
    k = norm(kab); m = KAB_PREFIX.match(kab); bare = norm(KAB_PREFIX.sub("", kab))
    if bare == norm(prov or "") or k in ("dki jakarta",): return prov_entry("kab_kota sama dengan provinsi")
    if m and m.group(1).lower().startswith("kota"): order = ["kota"]
    elif m: order = ["kab"]
    else: order = ["kab", "kota"]
    for t in order:
        hits = wd[t].get(bare, [])
        if len(hits) == 1:
            r = hits[0]; return {"lat": r["lat"], "lon": r["lon"], "tingkat": "kab_kota", "alasan": None, "sumber": src(r)}
        if len(hits) > 1:
            gagal.append({"provinsi": prov, "kab_kota": kab, "alasan": f"ambigu di Wikidata: {[h['qid'] for h in hits]}"}); return None
    gagal.append({"provinsi": prov, "kab_kota": kab, "alasan": "tidak ditemukan di Wikidata (kab/kota Indonesia dengan P625)"}); return None

def build_koordinat(insiden, kasus, wd_path):
    wd = load_wikidata(wd_path); gagal = []; out = {}
    pairs = sorted({(r.get("provinsi"), r.get("kab_kota")) for r in insiden + kasus}, key=lambda p: (p[0] or "", p[1] or ""))
    for prov, kab in pairs:
        e = resolve(prov, kab, wd, gagal)
        if e: out[f"{prov or ''}|{kab or ''}"] = e
    return out, gagal

# ---- validation
def validate(rows, id_col, url_col, sheet):
    problems, ok = [], []
    dup = {k for k, n in Counter(r.get(id_col) for r in rows).items() if n > 1}
    for r in rows:
        why = []
        if not r.get(id_col): why.append(f"{id_col} kosong")
        elif r[id_col] in dup: why.append(f"{id_col} ganda")
        u = r.get(url_col)
        if not u: why.append(f"tanpa tautan sumber ({url_col} kosong)")
        elif not URL_RE.match(str(u)): why.append(f"{url_col} bukan URL")
        if r.get("sumber_2_url") and not URL_RE.match(str(r["sumber_2_url"])): why.append("sumber_2_url bukan URL")
        if why: problems.append({"sheet": sheet, "row": r["_row"], "id": r.get(id_col), "alasan": why})
        else: ok.append(r)
    return ok, problems

def apply_titles(rows, url_col, titles):
    n = 0
    for r in rows:
        t = titles.get(r.get(url_col))
        if t and t.get("status") == "terverifikasi" and t.get("judul"):
            r["judul_sumber_1"] = t["judul"]; r["judul_status"] = "terverifikasi"; n += 1
    return n

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("xlsx", nargs="?", default=str(ROOT / "eppos-media-intake.xlsx"))
    ap.add_argument("--wikidata", default=str(ROOT / "scripts" / "wikidata_id_regions.json"))
    a = ap.parse_args()
    check_projection()
    wb = openpyxl.load_workbook(a.xlsx, data_only=True)
    periode = build_periode(wb)
    _, ins_raw = read_sheet(wb, "insiden"); _, kas_raw = read_sheet(wb, "kasus_resmi")
    ins, p1 = validate(ins_raw, "insiden_id", "sumber_1_url", "insiden")
    kas, p2 = validate(kas_raw, "kasus_id", "sumber_url", "kasus_resmi")
    titles = json.load(open(DATA / "judul_terverifikasi.json")) if (DATA / "judul_terverifikasi.json").exists() else {}
    nt = apply_titles(ins, "sumber_1_url", titles) + apply_titles(kas, "sumber_url", titles)
    insiden = [{c: r.get(c) for c in INSIDEN_COLS} for r in ins]
    kasus = [{c: r.get(c) for c in KASUS_COLS} for r in kas]
    koordinat, gagal = build_koordinat(insiden, kasus, a.wikidata)
    DATA.mkdir(exist_ok=True)
    def dump(name, obj): (DATA / name).write_text(json.dumps(obj, ensure_ascii=False, indent=1) + "\n")
    dump("periode.json", periode); dump("insiden.json", insiden); dump("kasus_resmi.json", kasus)
    dump("koordinat.json", koordinat); dump("koordinat_gagal.json", gagal); dump("validasi.json", {"ditolak": p1 + p2})
    keys = {f"{r['provinsi'] or ''}|{r['kab_kota'] or ''}" for r in insiden + kasus}
    manifest = {"version": dt.date.today().isoformat(), "generated_at": dt.datetime.now(dt.timezone(dt.timedelta(hours=7))).isoformat(timespec="seconds"),
                "xlsx": Path(a.xlsx).name, "xlsx_sha256": hashlib.sha256(open(a.xlsx, "rb").read()).hexdigest()[:16],
                "periode": len(periode), "insiden": len(insiden), "kasus_resmi": len(kasus),
                "insiden_ditolak": len(p1), "kasus_ditolak": len(p2), "judul_terverifikasi": nt,
                "koordinat": len(koordinat), "koordinat_gagal": len(gagal), "lokasi_dipakai": len(keys),
                "proyeksi": {"lonMin": LON_MIN, "lonMax": LON_MAX, "latMin": LAT_MIN, "latMax": LAT_MAX, "viewBox": "0 0 1000 420"}}
    dump("manifest.json", manifest)
    print(json.dumps(manifest, indent=1, ensure_ascii=False))
    for p in p1 + p2: print("  TOLAK", p)
    for g in gagal: print("  KOORDINAT", g)

if __name__ == "__main__":
    main()
