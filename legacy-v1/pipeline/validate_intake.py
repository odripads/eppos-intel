#!/usr/bin/env python3
"""Validate a filled `eppos-media-intake.xlsx` against the schema contract.

Usage:  python3 pipeline/validate_intake.py path/to/eppos-media-intake.xlsx [--json report.json]

Policy (spec §How I will work with you): rows missing a source URL are REJECTED, not repaired.
Every rejection names the sheet, the Excel row number and the reason. Warnings do not reject.
Exit code 1 if any row is rejected or the workbook headers do not match the contract.
"""
from __future__ import annotations
import argparse, json, re, sys, datetime as dt
from collections import Counter
import openpyxl
from schema import (INTAKE_SHEETS, DROPDOWNS, SOURCE_URL_COLUMN, PRIMARY_KEY, FOREIGN_KEYS,
                    DATE_COLUMNS, MONEY_COLUMNS, YEAR_COLUMNS, BARE_DOMAIN_RE, ISO_DATE_RE)

URL_RE = re.compile(r"^https?://[^\s]+$", re.I)
NAME_HINT_RE = re.compile(r"\b(pak|bu|bapak|ibu|sdr|sdri|saudara|kak|om|tante)\s+[A-Z]", re.I)
EXAMPLE_HINT = "contoh baris"
HEADER_ROW, FIRST_DATA_ROW = 2, 4   # row 1 = title, 2 = headers, 3 = hints, 4+ = data


def cell(v):
    if v is None: return ""
    if isinstance(v, (dt.date, dt.datetime)): return v.strftime("%Y-%m-%d")
    if isinstance(v, float) and v.is_integer(): return str(int(v))
    return str(v).strip()


def read_workbook(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    problems, tables = [], {}
    for sheet, cols in INTAKE_SHEETS.items():
        if sheet not in wb.sheetnames:
            problems.append({"sheet": sheet, "row": None, "level": "reject", "reason": "sheet hilang dari workbook"})
            continue
        ws = wb[sheet]
        headers = [cell(c.value) for c in ws[HEADER_ROW]][: len(cols)]
        if headers != cols:
            problems.append({"sheet": sheet, "row": HEADER_ROW, "level": "reject",
                             "reason": f"judul kolom berubah: {headers} != {cols}"})
            continue
        rows = []
        for r in ws.iter_rows(min_row=FIRST_DATA_ROW, values_only=False):
            values = [cell(c.value) for c in r][: len(cols)]
            if not any(values): continue
            rec = dict(zip(cols, values)); rec["_row"] = r[0].row
            if any(EXAMPLE_HINT in v.lower() or ".example" in v.lower() or re.search(r"\bcontoh\b", v, re.I) for v in values):
                rec["_example"] = True   # template example rows are skipped, never validated
            rows.append(rec)
        tables[sheet] = rows
    return tables, problems


def validate(tables):
    problems = []
    def rej(sheet, rec, reason): problems.append({"sheet": sheet, "row": rec["_row"], "level": "reject", "reason": reason})
    def warn(sheet, rec, reason): problems.append({"sheet": sheet, "row": rec["_row"], "level": "warn", "reason": reason})

    live = {s: [r for r in rows if not r.get("_example")] for s, rows in tables.items()}
    ids = {s: {r[PRIMARY_KEY[s]] for r in rows if r.get(PRIMARY_KEY[s])} for s, rows in live.items()}

    for sheet, rows in live.items():
        pk = PRIMARY_KEY[sheet]
        dup = {k for k, n in Counter(r[pk] for r in rows).items() if n > 1 and k}
        for rec in rows:
            if not rec[pk]: rej(sheet, rec, f"{pk} kosong")
            elif rec[pk] in dup: rej(sheet, rec, f"{pk} '{rec[pk]}' dipakai lebih dari satu baris")

            # 1. source URL — the non-negotiable
            col = SOURCE_URL_COLUMN[sheet]
            if col:
                val = rec.get(col, "")
                needs = not (sheet == "pejabat" and rec.get("punya_lhkpn") in ("tidak", "tidak diketahui"))
                if needs and not val: rej(sheet, rec, f"tanpa tautan sumber ({col} kosong)")
                elif val and not URL_RE.match(val): rej(sheet, rec, f"{col} bukan URL: '{val[:60]}'")
                elif not needs and not val: warn(sheet, rec, "tanpa LHKPN dan tanpa tautan sumber lain; masa jabatan tidak terverifikasi")
            for c in [c for c in rec if c.endswith("_url") and c != col]:
                if rec[c] and not URL_RE.match(rec[c]): rej(sheet, rec, f"{c} bukan URL: '{rec[c][:60]}'")

            # 2. dropdowns
            for (s, c), allowed in DROPDOWNS.items():
                if s == sheet and rec.get(c) and rec[c] not in allowed:
                    rej(sheet, rec, f"{c}='{rec[c]}' bukan pilihan yang diizinkan {allowed}")

            # 3. dates, money, years
            for c in DATE_COLUMNS.get(sheet, []):
                if rec.get(c) and not re.match(ISO_DATE_RE, rec[c]): rej(sheet, rec, f"{c}='{rec[c]}' bukan format YYYY-MM-DD")
            for c in MONEY_COLUMNS.get(sheet, []):
                if rec.get(c) and not re.match(r"^\d+$", rec[c]): rej(sheet, rec, f"{c}='{rec[c]}' harus angka saja")
            for c in YEAR_COLUMNS.get(sheet, []):
                if rec.get(c) and not re.match(r"^(20[12]\d)$", rec[c]): rej(sheet, rec, f"{c}='{rec[c]}' bukan tahun 2010–2029")

            # 4. per-sheet rules
            if sheet == "outlet":
                if rec["outlet_id"] and not re.match(BARE_DOMAIN_RE, rec["outlet_id"]):
                    rej(sheet, rec, f"outlet_id '{rec['outlet_id']}' bukan domain telanjang (huruf kecil, tanpa https://, tanpa www.)")
            if sheet == "relasi":
                s1, s2, st = rec.get("sumber_1_jenis"), rec.get("sumber_2_jenis"), rec.get("status")
                if st == "terkonfirmasi":
                    if not rec.get("sumber_2_url"): rej(sheet, rec, "status terkonfirmasi tanpa sumber_2_url (aturan dua sumber)")
                    elif s1 and s2 and s1 == s2: rej(sheet, rec, f"status terkonfirmasi tetapi kedua sumber sejenis ({s1}); harus berbeda jenis")
                if st == "dugaan" and not rec.get("catatan"): warn(sheet, rec, "status dugaan tanpa alasan di catatan")
                if not st: rej(sheet, rec, "status kosong")
            if sheet == "peristiwa":
                if not rec.get("bukti_jenis"): rej(sheet, rec, "bukti_jenis kosong; bukti harus menunjukkan peristiwanya terjadi")
                if not rec.get("tanggal"): rej(sheet, rec, "tanggal peristiwa kosong")
            if sheet == "info_a1":
                text = " ".join(rec.get(c, "") for c in ("peran_narasumber", "ringkasan"))
                if NAME_HINT_RE.search(text): rej(sheet, rec, "kemungkinan menyebut nama narasumber (sapaan + nama). Hapus namanya.")
            if sheet == "pejabat":
                if rec.get("mulai_menjabat") and rec.get("selesai_menjabat") and rec["selesai_menjabat"] < rec["mulai_menjabat"]:
                    rej(sheet, rec, "selesai_menjabat lebih awal dari mulai_menjabat")

    # 5. foreign keys
    for sheet, col, target in FOREIGN_KEYS:
        for rec in live.get(sheet, []):
            if rec.get(col) and rec[col] not in ids.get(target, set()):
                rej(sheet, rec, f"{col}='{rec[col]}' tidak ada di sheet {target}")
    return problems, live


def main():
    ap = argparse.ArgumentParser(); ap.add_argument("xlsx"); ap.add_argument("--json")
    a = ap.parse_args()
    tables, problems = read_workbook(a.xlsx)
    p2, live = validate(tables)
    problems += p2
    rejected = {(p["sheet"], p["row"]) for p in problems if p["level"] == "reject"}
    print(f"# Validasi {a.xlsx}")
    for s, rows in live.items():
        n_rej = sum(1 for r in rows if (s, r["_row"]) in rejected)
        n_ex = sum(1 for r in tables[s] if r.get("_example"))
        print(f"  {s:16s} {len(rows):4d} baris  ·  {len(rows)-n_rej:4d} diterima  ·  {n_rej:3d} ditolak  ·  {n_ex} baris contoh dilewati")
    for p in sorted(problems, key=lambda p: (p["level"] != "reject", p["sheet"], p["row"] or 0)):
        tag = "TOLAK" if p["level"] == "reject" else "peringatan"
        print(f"  [{tag}] {p['sheet']} baris {p['row']}: {p['reason']}")
    if a.json:
        json.dump({"problems": problems, "accepted": {s: [r for r in rows if (s, r["_row"]) not in rejected]
                   for s, rows in live.items()}}, open(a.json, "w"), ensure_ascii=False, indent=1)
    sys.exit(1 if rejected else 0)


if __name__ == "__main__":
    main()
